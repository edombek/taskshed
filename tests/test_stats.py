"""Statistics: aggregations, CSV/XLSX export, chart smoke tests."""

import os
from datetime import UTC, date, datetime, timedelta

import pytest

from taskshed.core.models import OccurrenceStatus, Task, TaskOccurrence
from taskshed.core.stats import per_day, per_place, per_task, trend
from taskshed.stats import charts, export

T0 = datetime(2026, 10, 1, tzinfo=UTC)


def mk(task_id, day, status):
    occ = TaskOccurrence(id=task_id * 100 + day, task_id=task_id,
                         due_at=T0 + timedelta(days=day),
                         reminder_at=T0, status=OccurrenceStatus(status))
    task = Task(id=task_id, user_id=1, place_id=1 if task_id == 1 else 2,
                title=f"T{task_id}")
    return occ, task


@pytest.fixture
def occs():
    return [
        mk(1, 0, "done"), mk(1, 1, "missed"), mk(1, 2, "done"),
        mk(2, 0, "done"), mk(2, 0, "pending"),
    ]


class TestAggregations:
    def test_per_task(self, occs):
        rows = per_task(occs)
        assert rows[0]["title"] == "T1" and rows[0]["percent"] == 67
        assert rows[1]["done"] == 1 and rows[1]["total"] == 2

    def test_per_day(self, occs):
        days = per_day(occs, "UTC")
        assert days[date(2026, 10, 1)] == {"done": 2, "missed": 0, "skipped": 0,
                                           "pending": 1}
        assert days[date(2026, 10, 2)]["missed"] == 1

    def test_per_place(self, occs):
        rows = per_place(occs, {1: "Дом", 2: "Работа"})
        by_name = {r["place"]: r for r in rows}
        assert by_name["Дом"]["total"] == 3
        assert by_name["Работа"]["done"] == 1

    def test_trend(self, occs):
        pts = trend(occs, "UTC")
        assert pts == [(date(2026, 9, 28), 3, 5)]  # one Monday-week bucket


class TestExport:
    def test_csv(self, tmp_path, occs):
        path = str(tmp_path / "s.csv")
        export.write_csv(per_task(occs), path)
        text = open(path, encoding="utf-8-sig").read()
        assert "T1;3;2;1;0;0;67" in text

    def test_xlsx(self, tmp_path, occs):
        from openpyxl import load_workbook

        path = str(tmp_path / "s.xlsx")
        export.write_xlsx(per_task(occs), per_place(occs, {1: "Home"}), path)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Sheet", "places"]
        assert wb.active.max_row == 3


class TestCharts:
    def test_daily_bar(self, tmp_path, occs):
        out = str(tmp_path / "d.png")
        charts.daily_bar(per_day(occs, "UTC"), "ru", out)
        assert os.path.getsize(out) > 500

    def test_top_tasks(self, tmp_path, occs):
        out = str(tmp_path / "t.png")
        charts.top_tasks(per_task(occs), "en", out)
        assert os.path.getsize(out) > 500

    def test_trend_line_and_heatmap(self, tmp_path, occs):
        out = str(tmp_path / "tr.png")
        charts.trend_line(trend(occs, "UTC"), "ru", out)
        assert os.path.getsize(out) > 500
        days = per_day(occs, "UTC")
        out2 = str(tmp_path / "h.png")
        assert charts.github_heatmap(days, "ru", out2) == out2
        assert os.path.getsize(out2) > 500
        # a year-long window stays within 53 week-columns
        far = {date(2025, 1, 1) + timedelta(days=i): {"done": 1} for i in range(400)}
        out3 = str(tmp_path / "h2.png")
        assert charts.github_heatmap(far, "en", out3) == out3

    def test_heatmap_window_at_least_one_month(self):
        # single recent day → still 5 week columns back from the anchor
        d0 = date(2026, 10, 26)
        first, n = charts._heatmap_window({d0: {"done": 1}}, end=d0,
                                          min_weeks=5, max_weeks=53)
        assert n == 5
        assert first <= d0 - timedelta(days=28)
        assert first.weekday() == 0  # window starts on Monday
        # sparse period a year long → clamped to 53 weeks
        sparse = {date(2025, 11, 1): {"done": 1}, d0: {"done": 2}}
        first, n = charts._heatmap_window(sparse, end=d0, min_weeks=5, max_weeks=53)
        assert n == 53
        # empty window helper must not be used for empty data
        assert charts.github_heatmap({}, "ru", "/tmp/none.png") == ""

    def test_heatmap_window_reaches_planned_days(self):
        # a task planned 6 days ahead must land inside the drawn window
        d0 = date(2026, 10, 26)
        days = {d0 - timedelta(days=3): {"done": 1},
                d0 + timedelta(days=6): {"pending": 1}}
        first, n = charts._heatmap_window(days, end=date(2026, 11, 5),
                                          min_weeks=5, max_weeks=53)
        last_monday = first + timedelta(weeks=n - 1)
        assert last_monday >= d0  # the week holding the planned day is drawn

    def test_hourly_heatmap(self, tmp_path):
        grid = {(date(2026, 10, 26), h): (h % 5) for h in range(24)}
        grid[(date(2026, 10, 27), 9)] = 3
        out = str(tmp_path / "hours.png")
        assert charts.hourly_heatmap(grid, "ru", out) == out
        assert os.path.getsize(out) > 500
        assert charts.hourly_heatmap({}, "ru", out) == ""

    def test_completions_by_hour(self):
        from taskshed.core.stats import completions_by_hour

        stamps = [datetime(2026, 10, 26, 7, 30, tzinfo=UTC),
                  datetime(2026, 10, 26, 7, 55, tzinfo=UTC),
                  datetime(2026, 10, 27, 20, 5, tzinfo=UTC)]
        grid = completions_by_hour(stamps, "Europe/Moscow")
        assert grid[(date(2026, 10, 26), 10)] == 2  # 07:30Z = 10:30 MSK
        assert grid[(date(2026, 10, 27), 23)] == 1  # 20:05Z = 23:05 MSK
